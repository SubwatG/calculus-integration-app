import os
import sys
import subprocess
import tempfile
import io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------
# 1. Authentic LaTeX to High-DPI Transparent PNG Generator
# ---------------------------------------------------------
def render_latex_to_png(latex_code: str, dpi: int = 300, color_hex: str = "0A2540") -> str:
    """Compiles a LaTeX math equation into a high-DPI transparent PNG using pdflatex + pdfcrop + pdftoppm."""
    td = tempfile.mkdtemp(prefix="latex_eq_")
    tex_path = os.path.join(td, "eq.tex")
    pdf_path = os.path.join(td, "eq.pdf")
    crop_path = os.path.join(td, "crop.pdf")
    out_prefix = os.path.join(td, "out")
    final_png = os.path.join(td, "out.png")

    tex_content = r"""\documentclass[24pt]{article}
\usepackage{amsmath,amssymb,amsfonts}
\usepackage{xcolor}
\definecolor{eqcolor}{HTML}{""" + color_hex + r"""}
\pagestyle{empty}
\begin{document}
\color{eqcolor}
\Huge
\[ """ + latex_code + r""" \]
\end{document}
"""
    with open(tex_path, "w", encoding="utf-8") as f:
        f.write(tex_content)

    # Compile TeX to PDF
    subprocess.run(["pdflatex", "-interaction=nonstopmode", "eq.tex"], cwd=td, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    if not os.path.exists(pdf_path):
        return None

    # Crop whitespace around equation
    subprocess.run(["pdfcrop", "eq.pdf", "crop.pdf"], cwd=td, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    if not os.path.exists(crop_path):
        crop_path = pdf_path

    # Rasterize to 300 DPI PNG
    subprocess.run(["pdftoppm", "-png", "-r", str(dpi), "-singlefile", crop_path, out_prefix], cwd=td, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    if os.path.exists(final_png):
        # Convert white background to transparent using PIL
        from PIL import Image
        im = Image.open(final_png).convert("RGBA")
        datas = im.getdata()
        new_data = []
        for item in datas:
            # If pixel is pure white or near-white, make it transparent
            if item[0] > 245 and item[1] > 245 and item[2] > 245:
                new_data.append((255, 255, 255, 0))
            else:
                new_data.append(item)
        im.putdata(new_data)
        transparent_png = os.path.join(td, "transparent.png")
        im.save(transparent_png, "PNG")
        return transparent_png

    return None

# ---------------------------------------------------------
# 2. PPTX Presentation Builder & Modern Academic Styler
# ---------------------------------------------------------
def create_deck(output_pptx: str):
    prs = Presentation()
    # 16:9 Widescreen standard
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    C_NAVY = RGBColor(10, 37, 64)       # #0A2540
    C_BLUE = RGBColor(29, 78, 216)      # #1D4ED8
    C_KU_GREEN = RGBColor(0, 102, 68)   # #006644
    C_DARK = RGBColor(15, 23, 42)       # #0F172A
    C_MUTED = RGBColor(100, 116, 139)   # #64748B
    C_CARD_BG = RGBColor(248, 250, 252) # #F8FAFC
    C_CARD_BORDER = RGBColor(226, 232, 240)
    C_CRIMSON = RGBColor(220, 38, 38)
    C_TEAL = RGBColor(13, 148, 136)

    def add_header(slide, section_label: str, title: str, slide_num: int):
        # Top Header Bar
        hdr_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = hdr_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_sec = tf.paragraphs[0]
        p_sec.text = section_label.upper()
        p_sec.font.size = Pt(11)
        p_sec.font.bold = True
        p_sec.font.color.rgb = C_BLUE
        p_sec.space_after = Pt(3)

        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(21)
        p_title.font.bold = True
        p_title.font.color.rgb = C_NAVY

        # Slide Number Badge
        num_box = slide.shapes.add_textbox(Inches(11.8), Inches(0.4), Inches(1.0), Inches(0.4))
        ntf = num_box.text_frame
        np = ntf.paragraphs[0]
        np.text = f"{slide_num:02d} / 15"
        np.font.size = Pt(11)
        np.font.bold = True
        np.font.color.rgb = C_MUTED
        np.alignment = PP_ALIGN.RIGHT

    def add_card(slide, left, top, width, height, title=None, title_color=C_NAVY, bg_color=C_CARD_BG):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = C_CARD_BORDER
        card.line.width = Pt(1)

        if title:
            tbox = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), Inches(0.4))
            tf = tbox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(15)
            p.font.bold = True
            p.font.color.rgb = title_color

        return card

    # =========================================================
    # SLIDE 1: Title Slide
    # =========================================================
    s1 = prs.slides.add_slide(blank_layout)
    # Background Accent Top
    top_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = C_BLUE
    top_bar.line.fill.background()

    # Title Card
    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.733), Inches(2.2))
    ttf = t_box.text_frame
    ttf.word_wrap = True
    tp0 = ttf.paragraphs[0]
    tp0.text = "เอกสารนำเสนอโครงร่างโครงงานคณิตศาสตร์ประยุกต์ ประจำปีการศึกษา 2569"
    tp0.font.size = Pt(13)
    tp0.font.bold = True
    tp0.font.color.rgb = C_KU_GREEN
    tp0.space_after = Pt(8)

    tp1 = ttf.add_paragraph()
    tp1.text = "การพัฒนาเว็บแอปพลิเคชันช่วยเรียนคณิตศาสตร์\nสำหรับแคลคูลัสเบื้องต้น"
    tp1.font.size = Pt(28)
    tp1.font.bold = True
    tp1.font.color.rgb = C_NAVY
    tp1.space_after = Pt(6)

    tp2 = ttf.add_paragraph()
    tp2.text = "Development of a Math Tutor Web Application for Introductory Calculus"
    tp2.font.size = Pt(15)
    tp2.font.color.rgb = C_MUTED

    # Left Info Card: Department
    add_card(s1, Inches(0.8), Inches(3.7), Inches(5.6), Inches(2.9), title="🏛️ สังกัดภาควิชาและสถาบัน", title_color=C_BLUE)
    c1_box = s1.shapes.add_textbox(Inches(1.1), Inches(4.3), Inches(5.0), Inches(2.1))
    c1_tf = c1_box.text_frame
    c1_tf.word_wrap = True
    c1_p1 = c1_tf.paragraphs[0]
    c1_p1.text = "ภาควิชา: วิทยาการคำนวณและเทคโนโลยีดิจิทัล\nคณะ: ศิลปศาสตร์และวิทยาศาสตร์\nสถาบัน: มหาวิทยาลัยเกษตรศาสตร์ วิทยาเขตกำแพงแสน\n\nอาจารย์ที่ปรึกษาโครงงาน:\nอาจารย์กิตติพงษ์ ทรัพย์วัฒนชัย"
    c1_p1.font.size = Pt(13)
    c1_p1.font.color.rgb = C_DARK

    # Right Info Card: Students
    add_card(s1, Inches(6.8), Inches(3.7), Inches(5.7), Inches(2.9), title="👩‍🎓 คณะผู้จัดทำโครงงาน", title_color=C_NAVY)
    c2_box = s1.shapes.add_textbox(Inches(7.1), Inches(4.3), Inches(5.1), Inches(2.1))
    c2_tf = c2_box.text_frame
    c2_tf.word_wrap = True
    c2_p1 = c2_tf.paragraphs[0]
    c2_p1.text = "1. นางสาวขวัญเรือน สำเนียงดี\n    รหัสนิสิต: 6621602928 • สาขาคณิตศาสตร์ประยุกต์\n\n2. นางสาวเบญจกัฬยาณี พิทยากูล\n    รหัสนิสิต: 6621602979 • สาขาคณิตศาสตร์ประยุกต์"
    c2_p1.font.size = Pt(13)
    c2_p1.font.color.rgb = C_DARK

    s1.notes_slide.notes_text_frame.text = (
        "บทพูดสไลด์ 1 (หน้าปก):\n"
        "กราบเรียนท่านคณะกรรมการผู้ทรงคุณวุฒิทุกท่าน พวกเราตัวแทนนายนิสิตสาขาวิชาคณิตศาสตร์ประยุกต์ "
        "ขอเสนอโครงร่างโครงงานในหัวข้อ 'การพัฒนาเว็บแอปพลิเคชันช่วยเรียนคณิตศาสตร์สำหรับแคลคูลัสเบื้องต้น' "
        "โดยมีอาจารย์กิตติพงษ์ ทรัพย์วัฒนชัย เป็นอาจารย์ที่ปรึกษาโครงงานค่ะ"
    )

    # =========================================================
    # SLIDE 2: Problem & Research Gap
    # =========================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "1. ที่มาและความสำคัญ", "สภาพปัญหาการเรียนรู้และช่องว่างของสื่อในปัจจุบัน (Research Gap)", 2)
    add_card(s2, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2), title="⚠️ ปัญหาของผู้เรียน (Student Pain Points)", title_color=C_CRIMSON)
    b2_1 = s2.shapes.add_textbox(Inches(1.05), Inches(2.3), Inches(5.1), Inches(4.4))
    tf2_1 = b2_1.text_frame
    tf2_1.word_wrap = True
    tf2_1.paragraphs[0].text = (
        "• ความเป็นนามธรรมสูง:\n"
        "  หัวข้อลิมิต อนุพันธ์ และปริพันธ์ ต้องอาศัยการเชื่อมโยงมโนทัศน์เรขาคณิตเข้ากับพีชคณิต\n\n"
        "• การจำสูตรแต่ขาดความเข้าใจเงื่อนไข:\n"
        "  ผู้เรียนมักจำสูตรสำเร็จ แต่ไม่เข้าใจที่มาและเงื่อนไขการใช้ (เช่น เงื่อนไขความต่อเนื่อง, การเลือกคู่ By Parts)\n\n"
        "• สื่อการเรียนรู้ไม่ตรงหลักสูตร:\n"
        "  ตัวอย่างบนอินเทอร์เน็ตกระจัดกระจาย ขาดระบบฝึกฝนภาษาไทยที่สอดคล้องกับแคลคูลัสมหาวิทยาลัย"
    )
    tf2_1.paragraphs[0].font.size = Pt(13)
    tf2_1.paragraphs[0].font.color.rgb = C_DARK

    add_card(s2, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2), title="🔍 ข้อจำกัดของเครื่องมือเดิม (The Black-Box Problem)", title_color=C_BLUE)
    b2_2 = s2.shapes.add_textbox(Inches(7.05), Inches(2.3), Inches(5.2), Inches(4.4))
    tf2_2 = b2_2.text_frame
    tf2_2.word_wrap = True
    tf2_2.paragraphs[0].text = (
        "• เครื่องคำนวณสำเร็จรูปเป็น Black-Box:\n"
        "  WolframAlpha, Symbolab แสดงคำตอบสุดท้ายทันที ข้ามกระบวนการให้เหตุผลทางคณิตศาสตร์\n\n"
        "• AI ทั่วไป (ChatGPT/LLM) มีภาพหลอน:\n"
        "  มักคิดเลขหรืออ้างทฤษฎีผิดพลาดในขั้นตอนละเอียด ไม่สามารถใช้เป็นแหล่งอ้างอิงเดี่ยวได้\n\n"
        "• เป้าหมายโครงงานเรา:\n"
        "  สร้าง Interactive Math Tutor ที่แสดงวิธีทำโปร่งใส ตรวจสอบความถูกต้องด้วย SymPy 100% และจัดกระบวนการคิดแบบ Scaffolding"
    )
    tf2_2.paragraphs[0].font.size = Pt(13)
    tf2_2.paragraphs[0].font.color.rgb = C_DARK

    # =========================================================
    # SLIDE 3: Comparison Matrix
    # =========================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "1. ที่มาและความสำคัญ", "ตารางเปรียบเทียบข้อจำกัดของเครื่องมือเดิมกับระบบของเรา", 3)
    add_card(s3, Inches(0.8), Inches(1.7), Inches(11.733), Inches(5.2), title="📊 ตารางวิเคราะห์เปรียบเทียบ (Feature Matrix)", title_color=C_NAVY)
    m_box = s3.shapes.add_textbox(Inches(1.1), Inches(2.3), Inches(11.1), Inches(4.3))
    mtf = m_box.text_frame
    mtf.word_wrap = True
    mtf.paragraphs[0].text = (
        "มิติการเปรียบเทียบ                   WolframAlpha / Symbolab       ChatGPT / AI แชทบอท           เว็บแอปพลิเคชันของเรา (โครงงานนี้)\n"
        "──────────────────────────────────────────────────────────────────────────────────────────────────\n"
        "1. ภาษาและบริบทเนื้อหา               ภาษาอังกฤษเป็นหลัก             ภาษาไทย (มักแปลทับศัพท์ผิด)   ภาษาไทยตรงตามหลักสูตร มก. 100%\n"
        "2. ความโปร่งใสของวิธีทำ             ข้ามขั้นตอน / ต้องจ่ายเงินเพิ่ม  มีภาพหลอน (คิดเลขผิดบ่อย)      แสดงขั้นตอนละเอียด โปร่งใสทุกจุด\n"
        "3. การตรวจสอบความถูกต้อง            Black-Box ตรวจสอบโค้ดไม่ได้     ไม่มีการรับประกันผลลัพธ์        ใช้ SymPy Computer Algebra System\n"
        "4. กราฟิกจำลองการโต้ตอบ             กราฟนิ่ง ปรับ slider ไม่ได้     ไม่มีกราฟไดนามิก               Reactive Sliders ปรับพารามิเตอร์สด\n"
        "5. การดักจับจุดผิด (Trap Detection) ไม่มีระบบเตือนการเลือกผิดทาง   เฉลยทันทีโดยไม่เตือนกับดัก     แจ้งเตือนทันทีเมื่อเลือกวิธีที่ทำให้โจทย์ยากขึ้น"
    )
    mtf.paragraphs[0].font.size = Pt(12)
    mtf.paragraphs[0].font.name = "Courier New"
    mtf.paragraphs[0].font.color.rgb = C_DARK

    # =========================================================
    # SLIDE 4: Objectives
    # =========================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "2. วัตถุประสงค์ของโครงงาน", "วัตถุประสงค์หลักและเป้าหมายการพัฒนา 3 มิติ", 4)
    objs = [
        ("🎯 วัตถุประสงค์ที่ 1: การพัฒนาระบบ", "เพื่อออกแบบและพัฒนาเว็บแอปพลิเคชันช่วยเรียนคณิตศาสตร์เรื่องแคลคูลัสเบื้องต้น ด้วยภาษาไพทอนผ่านเฟรมเวิร์ก Streamlit ที่มีความลื่นไหล ตอบสนองไว และใช้งานได้ทั้งบนคอมพิวเตอร์และแท็บเล็ต", C_BLUE),
        ("⚙️ วัตถุประสงค์ที่ 2: กลไกคณิตศาสตร์และกราฟิก", "เพื่อสร้างระบบคำนวณและแสดงขั้นตอนวิธีทำทางคณิตศาสตร์อย่างเป็นขั้นตอน ด้วย SymPy ควบคู่กับการแสดงภาพจำลองเรขาคณิตเชิงโต้ตอบ (Interactive Visualizations) ใน 4 หัวข้อหลัก", C_KU_GREEN),
        ("📈 วัตถุประสงค์ที่ 3: การประเมินและทดสอบ", "เพื่อทดสอบความถูกต้องของโมดูลคำนวณ 100% และประเมินประสิทธิผลการเรียนรู้ ตลอดจนความพึงพอใจของนิสิตกลุ่มตัวอย่างที่ใช้งานระบบตามเกณฑ์มาตรฐานทางวิชาการ", C_NAVY)
    ]
    for i, (otitle, odesc, ocolor) in enumerate(objs):
        top_pos = Inches(1.8 + i * 1.7)
        add_card(s4, Inches(0.8), top_pos, Inches(11.733), Inches(1.45), title=otitle, title_color=ocolor)
        obx = s4.shapes.add_textbox(Inches(1.1), top_pos + Inches(0.55), Inches(11.1), Inches(0.8))
        otf = obx.text_frame
        otf.word_wrap = True
        op = otf.paragraphs[0]
        op.text = odesc
        op.font.size = Pt(13.5)
        op.font.color.rgb = C_DARK

    # =========================================================
    # SLIDE 5: Scope & Modules
    # =========================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "3. ขอบเขตของโครงงาน", "4 โมดูลการเรียนรู้หลักตามโครงสร้างวิชาแคลคูลัส 1", 5)
    modules = [
        ("โมดูล 1: ลิมิตและความต่อเนื่อง", "• ลิมิตข้างเดียว (ซ้าย-ขวา)\n• นิยามความต่อเนื่อง 3 เงื่อนไข\n• กราฟไดนามิกจำลอง x วิ่งเข้าใกล้ a", Inches(0.8), Inches(1.8)),
        ("โมดูล 2: อนุพันธ์และการประยุกต์", "• นิยามอนุพันธ์จากอัตราการเปลี่ยนแปลง\n• ความชันและสมการเส้นสัมผัสโค้ง\n• การหาจุดวิกฤตและค่าสุดขีด", Inches(6.8), Inches(1.8)),
        ("โมดูล 3: ผลบวกรีมันน์และปริพันธ์", "• ผลบวกรีมันน์ (ซ้าย, ขวา, จุดกึ่งกลาง)\n• พื้นที่ระหว่างเส้นโค้งสองฟังก์ชัน\n• ปริมาตรทรงตันการหมุนแบบจาน/วงแหวน", Inches(0.8), Inches(4.3)),
        ("โมดูล 4: เทคนิคการอินทิเกรต", "• การแทนค่าด้วยตัวแปร u (u-Substitution)\n• การอินทิเกรตทีละส่วน (By Parts)\n• กฎ LIATE และระบบเตือนกับดักความซับซ้อน", Inches(6.8), Inches(4.3))
    ]
    for mtitle, mdesc, mx, my in modules:
        add_card(s5, mx, my, Inches(5.7), Inches(2.2), title=mtitle, title_color=C_BLUE)
        mbx = s5.shapes.add_textbox(mx + Inches(0.25), my + Inches(0.6), Inches(5.2), Inches(1.5))
        mtf = mbx.text_frame
        mtf.word_wrap = True
        mp = mtf.paragraphs[0]
        mp.text = mdesc
        mp.font.size = Pt(13)
        mp.font.color.rgb = C_DARK

    # =========================================================
    # SLIDE 6: System Architecture
    # =========================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "4. สถาปัตยกรรมระบบ", "โครงสร้างเชิงระบบ 3 ชั้น (3-Tier Layered Architecture)", 6)
    layers = [
        ("1. Presentation Layer (ส่วนปฏิสัมพันธ์ผู้ใช้)", "• Streamlit Multipage Framework คอนฟิก Academic Light Theme\n• Reactive Sliders และ Formula Inputs ปลอด Emoji ตามมาตรฐานทางการ\n• รองรับ Responsive Design ทั้งบน Desktop และ Tablet", C_BLUE, Inches(1.8)),
        ("2. Symbolic & Verification Layer (เครื่องยนต์คำนวณ)", "• SymPy Computer Algebra System คำนวณค่าจริงระดับ Analytical 100%\n• ตัวแปลงนิพจน์อัจฉริยะ (Smart Parser) จัดการ xor (^) และ implicit multiplication\n• Verification Engine หาอนุพันธ์ย้อนกลับเพื่อตรวจทานผลลัพธ์", C_KU_GREEN, Inches(3.55)),
        ("3. Pedagogical Scaffolding Layer (การจัดนั่งร้านทางปัญญา)", "• Conditional Reveal: ซ่อนขั้นตอนวิธีทำจนกว่าผู้เรียนจะตัดสินใจเลือกตัวแปร\n• Trap Detection: ตรวจสอบความซับซ้อนของดีกรีสมการและขึ้นกล่องเตือนทันที\n• Delayed-Feedback: ฝึกให้นิสิตคิดก่อนเฉลยคำตอบสำเร็จรูป", C_NAVY, Inches(5.3))
    ]
    for ltitle, ldesc, lcolor, ltop in layers:
        add_card(s6, Inches(0.8), ltop, Inches(11.733), Inches(1.5), title=ltitle, title_color=lcolor)
        lbx = s6.shapes.add_textbox(Inches(1.1), ltop + Inches(0.55), Inches(11.1), Inches(0.85))
        ltf = lbx.text_frame
        ltf.word_wrap = True
        lp = ltf.paragraphs[0]
        lp.text = ldesc
        lp.font.size = Pt(12.5)
        lp.font.color.rgb = C_DARK

    # =========================================================
    # SLIDE 7: Module 1 - Limits
    # =========================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "5. รายละเอียดโมดูลที่ 1", "ลิมิตของฟังก์ชันและความต่อเนื่อง (Limits & Continuity)", 7)
    add_card(s7, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2), title="📖 ทฤษฎีและเงื่อนไขหลัก", title_color=C_BLUE)
    s7_b1 = s7.shapes.add_textbox(Inches(1.05), Inches(2.3), Inches(5.1), Inches(2.2))
    s7_b1.text_frame.word_wrap = True
    s7_b1.text_frame.paragraphs[0].text = (
        "• การมีอยู่ของลิมิต:\n"
        "  ลิมิตสองด้านจะหาค่าได้ ก็ต่อเมื่อ ลิมิตซ้ายและขวามีค่าเท่ากัน\n\n"
        "• เงื่อนไขความต่อเนื่อง 3 ประการ:\n"
        "  1. f(a) หาค่าได้ (มีนิยามในโดเมน)\n"
        "  2. lim f(x) เมื่อ x เข้าใกล้ a หาค่าได้\n"
        "  3. ค่าลิมิตเท่ากับค่าของฟังก์ชัน f(a)"
    )
    s7_b1.text_frame.paragraphs[0].font.size = Pt(13)
    s7_b1.text_frame.paragraphs[0].font.color.rgb = C_DARK

    # Render authentic LaTeX for Limits
    lim_png = render_latex_to_png(r"\lim_{x \to a^-} f(x) = \lim_{x \to a^+} f(x) = L \iff \lim_{x \to a} f(x) = L")
    if lim_png:
        s7.shapes.add_picture(lim_png, Inches(1.1), Inches(4.9), width=Inches(5.0))

    add_card(s7, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2), title="💡 การแสดงผลเชิงโต้ตอบในแอปเรา", title_color=C_NAVY)
    s7_b2 = s7.shapes.add_textbox(Inches(7.05), Inches(2.3), Inches(5.2), Inches(4.3))
    s7_b2.text_frame.word_wrap = True
    s7_b2.text_frame.paragraphs[0].text = (
        "• สไลเดอร์ปรับค่า x เข้าใกล้จุดวิกฤต:\n"
        "  ผู้เรียนสามารถเลื่อน slider เพื่อสังเกตค่า f(x) จากด้านซ้ายและขวาได้แบบ Real-time\n\n"
        "• ตรวจจับจุดไม่ต่อเนื่องอัตโนมัติ:\n"
        "  ระบบไฮไลต์จุดที่ฟังก์ชันขาดตอน พร้อมระบุประเภทความไม่ต่อเนื่อง (Removable, Jump, Infinite)\n\n"
        "• กราฟิก Plotly ควบคู่ตารางค่าตัวเลข:\n"
        "  เชื่อมโยงมโนทัศน์กราฟิกเรขาคณิตเข้ากับตารางตัวเลขเชิงตัวเลข ช่วยให้นิสิตเห็นภาพชัดเจน"
    )
    s7_b2.text_frame.paragraphs[0].font.size = Pt(13)
    s7_b2.text_frame.paragraphs[0].font.color.rgb = C_DARK

    s7.notes_slide.notes_text_frame.text = "LaTeX: \\lim_{x \\to a^-} f(x) = \\lim_{x \\to a^+} f(x) = L \\iff \\lim_{x \\to a} f(x) = L"

    # =========================================================
    # SLIDE 8: Module 2 - Derivatives
    # =========================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "5. รายละเอียดโมดูลที่ 2", "อนุพันธ์และการประยุกต์ (Derivatives & Applications)", 8)
    add_card(s8, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2), title="📖 ทฤษฎีและสมการหลัก", title_color=C_BLUE)
    s8_b1 = s8.shapes.add_textbox(Inches(1.05), Inches(2.3), Inches(5.1), Inches(2.0))
    s8_b1.text_frame.word_wrap = True
    s8_b1.text_frame.paragraphs[0].text = (
        "• นิยามอนุพันธ์จากลิมิต:\n"
        "  อัตราการเปลี่ยนแปลงขณะใดขณะหนึ่ง หรือความชันของเส้นสัมผัสกราฟ ณ จุด x\n\n"
        "• สมการเส้นสัมผัสโค้ง:\n"
        "  ความชัน m = f'(x_0) และสมการเส้นสัมผัส y - y_0 = m(x - x_0)"
    )
    s8_b1.text_frame.paragraphs[0].font.size = Pt(13)
    s8_b1.text_frame.paragraphs[0].font.color.rgb = C_DARK

    # Render authentic LaTeX for Derivative definition
    deriv_png = render_latex_to_png(r"f'(x) = \lim_{h \to 0} \frac{f(x+h) - f(x)}{h}")
    if deriv_png:
        s8.shapes.add_picture(deriv_png, Inches(1.2), Inches(4.8), width=Inches(4.8))

    add_card(s8, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2), title="💡 การแสดงผลเชิงโต้ตอบในแอปเรา", title_color=C_NAVY)
    s8_b2 = s8.shapes.add_textbox(Inches(7.05), Inches(2.3), Inches(5.2), Inches(4.3))
    s8_b2.text_frame.word_wrap = True
    s8_b2.text_frame.paragraphs[0].text = (
        "• แอนิเมชันเส้นตัดกราฟ (Secant Line) ลู่เข้าสู่เส้นสัมผัส:\n"
        "  ปรับค่า h ให้เล็กลงเรื่อยๆ นิสิตจะเห็นเส้น secant หมุนเข้าทาบเส้น tangent พอดี\n\n"
        "• การค้นหาจุดวิกฤตและทดสอบค่าสุดขีด:\n"
        "  ระบบแก้สมการ f'(x) = 0 ด้วย SymPy แสดงจุดสูงสุด/ต่ำสุดสัมพัทธ์บนกราฟทันที\n\n"
        "• แดชบอร์ดตรวจสอบกฎการดิฟ:\n"
        "  รองรับ Product Rule, Quotient Rule และ Chain Rule พร้อมวิธีทำแยกพจน์"
    )
    s8_b2.text_frame.paragraphs[0].font.size = Pt(13)
    s8_b2.text_frame.paragraphs[0].font.color.rgb = C_DARK

    s8.notes_slide.notes_text_frame.text = "LaTeX: f'(x) = \\lim_{h \\to 0} \\frac{f(x+h) - f(x)}{h}"

    # =========================================================
    # SLIDE 9: Module 3 - Riemann Sums
    # =========================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "5. รายละเอียดโมดูลที่ 3", "ผลบวกรีมันน์และปริพันธ์จำกัดเขต (Riemann Sums)", 9)
    add_card(s9, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2), title="📖 ทฤษฎีผลบวกรีมันน์", title_color=C_BLUE)
    s9_b1 = s9.shapes.add_textbox(Inches(1.05), Inches(2.3), Inches(5.1), Inches(1.8))
    s9_b1.text_frame.word_wrap = True
    s9_b1.text_frame.paragraphs[0].text = (
        "• การแบ่งช่วงย่อยและการประมาณพื้นที่:\n"
        "  แบ่งช่วง [a, b] ออกเป็น n ส่วนย่อย ความกว้างแต่ละแท่งเท่ากับ Δx = (b-a)/n\n\n"
        "• ลิมิตของผลบวกรีมันน์สู่ปริพันธ์จำกัดเขต:\n"
        "  เมื่อ n เข้าสู่อนันต์ ผลรวมของแท่งสี่เหลี่ยมจะเท่ากับพื้นที่ใต้กราฟแท้จริง"
    )
    s9_b1.text_frame.paragraphs[0].font.size = Pt(13)
    s9_b1.text_frame.paragraphs[0].font.color.rgb = C_DARK

    # Render authentic LaTeX for Riemann Sum
    riemann_png = render_latex_to_png(r"\int_a^b f(x) \, dx = \lim_{n \to \infty} \sum_{i=1}^n f(x_i^*) \Delta x")
    if riemann_png:
        s9.shapes.add_picture(riemann_png, Inches(1.0), Inches(4.7), width=Inches(5.2))

    add_card(s9, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2), title="💡 การแสดงผลเชิงโต้ตอบในแอปเรา", title_color=C_NAVY)
    s9_b2 = s9.shapes.add_textbox(Inches(7.05), Inches(2.3), Inches(5.2), Inches(4.3))
    s9_b2.text_frame.word_wrap = True
    s9_b2.text_frame.paragraphs[0].text = (
        "• Slider ปรับจำนวนช่วง n (1 ถึง 100 ช่วง):\n"
        "  นิสิตเลื่อนแถบเพื่อดูแท่งสี่เหลี่ยมซอยถี่ขึ้นเรื่อยๆ จนเต็มพื้นที่ใต้เส้นโค้ง\n\n"
        "• เปรียบเทียบ 3 วิธีคำนวณ:\n"
        "  1. Left Riemann Sum (จุดปลายซ้าย)\n"
        "  2. Right Riemann Sum (จุดปลายขวา)\n"
        "  3. Midpoint Rule (จุดกึ่งกลางช่วง)\n\n"
        "• แสดงตารางคำนวณค่าคลาดเคลื่อน (Error Analysis):\n"
        "  เทียบผลบวกกับค่าอินทิเกรตจริง เพื่อให้นิสิตเข้าใจอัตราการลู่เข้าของแต่ละวิธี"
    )
    s9_b2.text_frame.paragraphs[0].font.size = Pt(13)
    s9_b2.text_frame.paragraphs[0].font.color.rgb = C_DARK

    s9.notes_slide.notes_text_frame.text = "LaTeX: \\int_a^b f(x) \\, dx = \\lim_{n \\to \\infty} \\sum_{i=1}^n f(x_i^*) \\Delta x"

    # =========================================================
    # SLIDE 10: Module 3 (Cont.) - Area & Volume
    # =========================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "5. รายละเอียดโมดูลที่ 3 (ต่อ)", "การประยุกต์ปริพันธ์: พื้นที่และปริมาตรทรงตันการหมุน", 10)
    add_card(s10, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2), title="📖 สูตรคำนวณพื้นที่และปริมาตร", title_color=C_BLUE)
    s10_b1 = s10.shapes.add_textbox(Inches(1.05), Inches(2.3), Inches(5.1), Inches(1.2))
    s10_b1.text_frame.word_wrap = True
    s10_b1.text_frame.paragraphs[0].text = "• พื้นที่ระหว่างเส้นโค้งสองเส้น f(x) และ g(x):"
    s10_b1.text_frame.paragraphs[0].font.size = Pt(13)
    s10_b1.text_frame.paragraphs[0].font.color.rgb = C_DARK

    # Render authentic LaTeX for Area
    area_png = render_latex_to_png(r"A = \int_a^b \big| f(x) - g(x) \big| \, dx")
    if area_png:
        s10.shapes.add_picture(area_png, Inches(1.2), Inches(3.0), width=Inches(4.6))

    s10_b1_2 = s10.shapes.add_textbox(Inches(1.05), Inches(4.3), Inches(5.1), Inches(0.8))
    s10_b1_2.text_frame.word_wrap = True
    s10_b1_2.text_frame.paragraphs[0].text = "• ปริมาตรทรงตันการหมุนรอบแกน (Washer Method):"
    s10_b1_2.text_frame.paragraphs[0].font.size = Pt(13)
    s10_b1_2.text_frame.paragraphs[0].font.color.rgb = C_DARK

    # Render authentic LaTeX for Volume
    vol_png = render_latex_to_png(r"V = \pi \int_a^b \Big( [R(x)]^2 - [r(x)]^2 \Big) \, dx")
    if vol_png:
        s10.shapes.add_picture(vol_png, Inches(0.9), Inches(5.0), width=Inches(5.3))

    add_card(s10, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2), title="💡 การแสดงผลเชิงโต้ตอบในแอปเรา", title_color=C_NAVY)
    s10_b2 = s10.shapes.add_textbox(Inches(7.05), Inches(2.3), Inches(5.2), Inches(4.3))
    s10_b2.text_frame.word_wrap = True
    s10_b2.text_frame.paragraphs[0].text = (
        "• ไฮไลต์พื้นที่แรงเงาระหว่างเส้นโค้งอัตโนมัติ:\n"
        "  ระบบคำนวณจุดตัดของกราฟทั้งสองเส้น และแรเงาพื้นที่ส่วนที่ล้อมรอบให้เห็นชัดเจน\n\n"
        "• การจำลองการหมุน 3 มิติ (3D Solid of Revolution):\n"
        "  หมุนเส้นโค้งรอบแกน X หรือแกน Y เพื่อสร้างรูปทรงตัน 3D แบบ Interactive Plotly\n\n"
        "• รองรับทั้ง Disk Method, Washer Method และ Shell Method:\n"
        "  แสดงรัศมีภายนอก R(x) และรัศมีภายใน r(x) พร้อมเทียบวิธีตั้งสมการ"
    )
    s10_b2.text_frame.paragraphs[0].font.size = Pt(13)
    s10_b2.text_frame.paragraphs[0].font.color.rgb = C_DARK

    s10.notes_slide.notes_text_frame.text = (
        "LaTeX Area: A = \\int_a^b |f(x) - g(x)| \\, dx\n"
        "LaTeX Volume: V = \\pi \\int_a^b ([R(x)]^2 - [r(x)]^2) \\, dx"
    )

    # =========================================================
    # SLIDE 11: Module 4 - Integration by Parts
    # =========================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "5. รายละเอียดโมดูลที่ 4", "เทคนิค By Parts และกฎการตัดสินใจ LIATE", 11)
    add_card(s11, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2), title="📖 สูตรหลัก By Parts และกฎ LIATE", title_color=C_BLUE)
    
    # Render authentic LaTeX for By Parts
    byparts_png = render_latex_to_png(r"\int u \, dv = uv - \int v \, du")
    if byparts_png:
        s11.shapes.add_picture(byparts_png, Inches(1.2), Inches(2.3), width=Inches(4.8))

    s11_b1 = s11.shapes.add_textbox(Inches(1.05), Inches(3.6), Inches(5.1), Inches(3.1))
    s11_b1.text_frame.word_wrap = True
    s11_b1.text_frame.paragraphs[0].text = (
        "กฎลำดับความสำคัญ LIATE สำหรับเลือกตัวแปร u:\n"
        "  1. L - Logarithmic (ลอการิทึม เช่น ln x) → เลือกเป็น u เสมอ\n"
        "  2. I - Inverse Trig (ตรีโกณผกผัน เช่น arctan x) → ควรเป็น u\n"
        "  3. A - Algebraic (พีชคณิต เช่น x, x²) → ดีกรีจะลดลงเมื่อดิฟ\n"
        "  4. T - Trigonometric (ตรีโกณ เช่น sin x, cos x) → มักเป็น dv\n"
        "  5. E - Exponential (เอกซ์โพเนนเชียล เช่น eˣ) → เลือกเป็น dv เสมอ"
    )
    s11_b1.text_frame.paragraphs[0].font.size = Pt(12)
    s11_b1.text_frame.paragraphs[0].font.color.rgb = C_DARK

    add_card(s11, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2), title="⚖️ กรณีศึกษา: ทางเลือกถูกต้อง VS กับดัก", title_color=C_NAVY)
    s11_b2 = s11.shapes.add_textbox(Inches(7.05), Inches(2.3), Inches(5.2), Inches(4.3))
    s11_b2.text_frame.word_wrap = True
    s11_b2.text_frame.paragraphs[0].text = (
        "โจทย์ตัวอย่าง: ∫ x eˣ dx\n"
        "───────────────────────────────────────\n"
        "✅ ทางเลือกที่ถูกต้องตามกฎ LIATE:\n"
        "  ให้ u = x และ dv = eˣ dx\n"
        "  จะได้ du = dx และ v = eˣ\n"
        "  สูตร: ∫ x eˣ dx = x eˣ - ∫ eˣ dx = (x - 1)eˣ + C\n"
        "  (พจน์ใหม่ ∫ eˣ dx ดีกรีลดลง อินทิเกรตจบได้ทันที)\n\n"
        "❌ กับดักความซับซ้อน (Trap Detected):\n"
        "  หากสลับให้ u = eˣ และ dv = x dx\n"
        "  จะได้ du = eˣ dx และ v = x²/2\n"
        "  สูตร: ∫ x eˣ dx = (x²/2)eˣ - ∫ (x²/2)eˣ dx\n"
        "  ⚠️ ดีกรีของ x บวมขึ้นจาก 1 เป็น 2 โจทย์ยากกว่าเดิม!"
    )
    s11_b2.text_frame.paragraphs[0].font.size = Pt(12.5)
    s11_b2.text_frame.paragraphs[0].font.color.rgb = C_DARK

    s11.notes_slide.notes_text_frame.text = "LaTeX: \\int u \\, dv = uv - \\int v \\, du"

    # =========================================================
    # SLIDE 12: Pedagogical Scaffolding
    # =========================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "6. กลไกการเรียนรู้", "การจัดนั่งร้านทางปัญญาและดักจับจุดผิด (Pedagogical Scaffolding)", 12)
    peds = [
        ("1. การเปิดเผยตามเงื่อนไข (Conditional Reveal)", "ระบบจะไม่รีบเฉลยวิธีทำทั้งหมดในทันที แต่จะบังคับให้นิสิตตัดสินใจเลือกคู่ตัวแปร (เช่น เลือก u และ dv) ด้วยตนเองก่อน เพื่อสร้างกระบวนการคิดเชิงวิเคราะห์ (Active Engagement)", C_BLUE, Inches(1.8)),
        ("2. ระบบตรวจจับกับดักความซับซ้อน (Complexity Trap Detection)", "หากนิสิตเลือกตัวแปรผิดทาง ระบบจะไม่เพียงแค่บอกว่า 'ผิด' แต่จะแสดงขั้นตอนให้เห็นชัดเจนว่าทำไมพจน์ใหม่ถึงซับซ้อนกว่าโจทย์เดิม และแนะนำกฎ LIATE เพื่อแก้ไขทิศทาง", C_CRIMSON, Inches(3.55)),
        ("3. การพิสูจน์ย้อนกลับด้วย SymPy (Mathematical Verification)", "ทุกคำตอบสุดท้าย ระบบจะหาอนุพันธ์ (d/dx) ย้อนกลับเพื่อพิสูจน์ให้ผู้เรียนเห็นจริงว่าได้ฟังก์ชันโจทย์ตั้งต้น 100% สอดคล้องกับ Fundamental Theorem of Calculus", C_KU_GREEN, Inches(5.3))
    ]
    for ptitle, pdesc, pcolor, ptop in peds:
        add_card(s12, Inches(0.8), ptop, Inches(11.733), Inches(1.5), title=ptitle, title_color=pcolor)
        pbx = s12.shapes.add_textbox(Inches(1.1), ptop + Inches(0.55), Inches(11.1), Inches(0.85))
        ptf = pbx.text_frame
        ptf.word_wrap = True
        pp = ptf.paragraphs[0]
        pp.text = pdesc
        pp.font.size = Pt(13)
        pp.font.color.rgb = C_DARK

    # =========================================================
    # SLIDE 13: Methodology & Plan
    # =========================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "7. ระเบียบวิธีวิจัยและแผนงาน", "แผนการดำเนินงาน 6 ขั้นตอนตามหลักการวิศวกรรมซอฟต์แวร์", 13)
    add_card(s13, Inches(0.8), Inches(1.7), Inches(11.733), Inches(5.2), title="🗓️ ตารางขั้นตอนการดำเนินงานโครงงาน", title_color=C_NAVY)
    plan_box = s13.shapes.add_textbox(Inches(1.1), Inches(2.3), Inches(11.1), Inches(4.3))
    ptf = plan_box.text_frame
    ptf.word_wrap = True
    ptf.paragraphs[0].text = (
        "ขั้นตอนการดำเนินงาน                          ระยะเวลาดำเนินการ        ผลผลิตที่ได้ (Deliverables)\n"
        "──────────────────────────────────────────────────────────────────────────────────────────────────\n"
        "1. ทบทวนวรรณกรรมและหลักสูตรแคลคูลัส 1        เดือนที่ 1 - 2           เอกสารข้อกำหนดความต้องการและขอบเขตโจทย์\n"
        "2. ออกแบบโครงสร้างระบบและ UI Wireframe      เดือนที่ 2 - 3           ผังระบบ 3-Tier และม็อกอัปหน้าจอ Streamlit\n"
        "3. พัฒนาโมดูลคำนวณ SymPy และ Reactive UI      เดือนที่ 3 - 5           โค้ดโปรแกรม 4 โมดูลหลักพร้อมทำงาน\n"
        "4. ทดสอบความถูกต้องทางคณิตศาสตร์ (Unit Test)  เดือนที่ 5 - 6           ชุดทดสอบ PyTest ครอบคลุมทุกเงื่อนไขโจทย์ 100%\n"
        "5. ทดลองใช้กับกลุ่มตัวอย่างนิสิตและเก็บข้อมูล      เดือนที่ 6 - 7           คะแนนแบบทดสอบและผลประเมินความพึงพอใจ\n"
        "6. วิเคราะห์ผลทางสถิติและจัดทำรายงานฉบับสมบูรณ์   เดือนที่ 7 - 8           เล่มรายงานโครงงานฉบับสมบูรณ์และบทความวิชาการ"
    )
    ptf.paragraphs[0].font.size = Pt(12)
    ptf.paragraphs[0].font.name = "Courier New"
    ptf.paragraphs[0].font.color.rgb = C_DARK

    # =========================================================
    # SLIDE 14: Expected Outcomes
    # =========================================================
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "8. ประโยชน์ที่คาดว่าจะได้รับ", "คุณค่าของผลผลิตโครงงานต่อผู้เรียน อาจารย์ และงานวิจัย", 14)
    add_card(s14, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2), title="🌟 ประโยชน์ต่อผู้เรียนและอาจารย์", title_color=C_KU_GREEN)
    s14_b1 = s14.shapes.add_textbox(Inches(1.05), Inches(2.3), Inches(5.1), Inches(4.3))
    s14_b1.text_frame.word_wrap = True
    s14_b1.text_frame.paragraphs[0].text = (
        "• ประโยชน์ต่อนิสิตผู้เรียน:\n"
        "  - ลดอัตราความเข้าใจผิด (Misconceptions) ในหัวข้อแคลคูลัสเบื้องต้น\n"
        "  - ฝึกกระบวนการคิดและแก้ปัญหาเป็นลำดับขั้น ไม่พึ่งพาการลอกคำตอบ\n"
        "  - เข้าถึงสื่อการสอนคุณภาพสูงภาษาไทยได้ฟรีตลอด 24 ชั่วโมง\n\n"
        "• ประโยชน์ต่ออาจารย์ผู้สอน:\n"
        "  - ใช้เป็นสื่อสาธิตมโนทัศน์เรขาคณิตในห้องเรียน Active Learning\n"
        "  - มีคลังโจทย์และกราฟิกจำลองพร้อมใช้ในการบรรยายและตรวจข้อสอบ"
    )
    s14_b1.text_frame.paragraphs[0].font.size = Pt(13)
    s14_b1.text_frame.paragraphs[0].font.color.rgb = C_DARK

    add_card(s14, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2), title="📊 แผนการประเมินผลเชิงประจักษ์", title_color=C_BLUE)
    s14_b2 = s14.shapes.add_textbox(Inches(7.05), Inches(2.3), Inches(5.2), Inches(4.3))
    s14_b2.text_frame.word_wrap = True
    s14_b2.text_frame.paragraphs[0].text = (
        "• การทดสอบผลสัมฤทธิ์ทางการเรียน (Learning Gains):\n"
        "  ทดสอบกลุ่มตัวอย่างนิสิตด้วย Pre-test และ Post-test และวิเคราะห์ผลต่างด้วยสถิติ Dependent Samples t-test\n\n"
        "• การประเมินความพึงพอใจและการใช้งาน (SUS):\n"
        "  ใช้แบบสอบถาม System Usability Scale (SUS) 10 ข้อ เพื่อวัดระดับความง่ายในการใช้งานและความพึงพอใจของผู้เรียน\n\n"
        "• การตรวจสอบความถูกต้องของระบบ (Mathematical Rigor):\n"
        "  รัน Automated Test Suites เทียบผลเฉลยของ SymPy กับเฉลยตำรามาตรฐานมากกว่า 100 ข้อสอบ"
    )
    s14_b2.text_frame.paragraphs[0].font.size = Pt(13)
    s14_b2.text_frame.paragraphs[0].font.color.rgb = C_DARK

    # =========================================================
    # SLIDE 15: Conclusion & Q&A
    # =========================================================
    s15 = prs.slides.add_slide(blank_layout)
    top_bar15 = s15.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    top_bar15.fill.solid()
    top_bar15.fill.fore_color.rgb = C_KU_GREEN
    top_bar15.line.fill.background()

    add_card(s15, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.1), title="🎓 สรุปภาพรวมและเปิดรับข้อเสนอแนะ (Q&A)", title_color=C_NAVY)
    s15_box = s15.shapes.add_textbox(Inches(2.0), Inches(2.0), Inches(9.333), Inches(4.0))
    s15_tf = s15_box.text_frame
    s15_tf.word_wrap = True
    s15_p = s15_tf.paragraphs[0]
    s15_p.text = (
        "สรุปสาระสำคัญของโครงงาน:\n"
        "• เปลี่ยนผ่านจากเครื่องมือคำนวณแบบ Black-Box สู่ 'ผู้ช่วยสอนอัจฉริยะที่โปร่งใสและโต้ตอบได้'\n"
        "• ผสานพลังของ Streamlit (UI ตอบสนองไว) + SymPy (คณิตศาสตร์แม่นยำ 100%)\n"
        "• จัดกระบวนการเรียนรู้แบบ Scaffolding และดักจับกับดักความซับซ้อนตามหลักการสอนสากล\n\n"
        "ข้อมูลการพัฒนาและ Source Code:\n"
        "• GitHub Repository: SubwatG/calculus-integration-app\n"
        "• สถาบัน: ภาควิชาวิทยาการคำนวณและเทคโนโลยีดิจิทัล คณะศิลปศาสตร์และวิทยาศาสตร์ มหาวิทยาลัยเกษตรศาสตร์\n\n"
        "ขอขอบพระคุณคณะกรรมการทุกท่านเป็นอย่างสูง\n"
        "พร้อมรับฟังข้อเสนอแนะและตอบข้อซักถามค่ะ"
    )
    s15_p.font.size = Pt(14)
    s15_p.font.color.rgb = C_DARK

    s15.notes_slide.notes_text_frame.text = (
        "บทพูดสไลด์ 15 (บทสรุป):\n"
        "ทั้งหมดนี้คือโครงร่างการพัฒนาเว็บแอปพลิเคชันช่วยเรียนคณิตศาสตร์สำหรับแคลคูลัสเบื้องต้น "
        "ที่มุ่งหวังจะแก้ปัญหาการเรียนรู้ของนิสิตอย่างเป็นรูปธรรม พวกเราขอขอบพระคุณท่านคณะกรรมการทุกท่าน "
        "และพร้อมรับฟังทุกคำชี้แนะเพื่อนำไปพัฒนาโครงงานให้สมบูรณ์ยิ่งขึ้นค่ะ"
    )

    # Save to disk
    prs.save(output_pptx)
    print(f"✅ Successfully generated professional LaTeX-enhanced deck: {output_pptx}")

if __name__ == "__main__":
    out_file = sys.argv[1] if len(sys.argv) > 1 else "/home/kitti/Downloads/calculus-tutor-presentation.pptx"
    create_deck(out_file)
